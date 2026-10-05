import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def create_report():
    doc = docx.Document()

    # Set page margins
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    primary_color = RGBColor(124, 76, 38)     # #7C4C26 - Kanaga Brown
    secondary_color = RGBColor(184, 134, 11) # Dark goldenrod
    dark_text = RGBColor(45, 26, 10)         # Very dark warm brown
    muted_text = RGBColor(100, 100, 100)

    # Document Title
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_title = p_title.add_run("KANAGA CONSULTING")
    run_title.font.name = "Arial"
    run_title.font.size = Pt(22)
    run_title.font.bold = True
    run_title.font.color.rgb = primary_color

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_sub = p_sub.add_run("Portail d'Évaluation de la Performance\n")
    run_sub.font.name = "Arial"
    run_sub.font.size = Pt(14)
    run_sub.font.bold = True
    run_sub.font.color.rgb = RGBColor(90, 90, 90)

    run_sub2 = p_sub.add_run("Rapport Complet : Processus Métier & Points Fondamentaux des Fonctionnalités d'Évaluation")
    run_sub2.font.name = "Arial"
    run_sub2.font.size = Pt(12)
    run_sub2.font.italic = True
    run_sub2.font.color.rgb = muted_text

    p_divider = doc.add_paragraph()
    p_divider.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_div = p_divider.add_run("―" * 45)
    run_div.font.color.rgb = primary_color

    # Helper for adding styled headings
    def add_h1(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(18)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.font.name = "Arial"
        run.font.size = Pt(15)
        run.font.bold = True
        run.font.color.rgb = primary_color
        return p

    def add_h2(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.font.name = "Arial"
        run.font.size = Pt(12.5)
        run.font.bold = True
        run.font.color.rgb = RGBColor(70, 70, 70)
        return p

    def add_body(text, bold_prefix="", italic=False):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.15
        if bold_prefix:
            r_bold = p.add_run(bold_prefix)
            r_bold.font.name = "Arial"
            r_bold.font.size = Pt(10.5)
            r_bold.font.bold = True
            r_bold.font.color.rgb = dark_text
        r = p.add_run(text)
        r.font.name = "Arial"
        r.font.size = Pt(10.5)
        r.font.italic = italic
        r.font.color.rgb = dark_text
        return p

    def add_bullet(bold_prefix, text):
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.15
        if bold_prefix:
            r_bold = p.add_run(bold_prefix)
            r_bold.font.name = "Arial"
            r_bold.font.size = Pt(10.5)
            r_bold.font.bold = True
            r_bold.font.color.rgb = dark_text
        r = p.add_run(text)
        r.font.name = "Arial"
        r.font.size = Pt(10.5)
        r.font.color.rgb = dark_text
        return p

    # 1. INTRODUCTION & VISION GLOBALE
    add_h1("1. Introduction & Vision Globale")
    add_body("Le module d'évaluation du Portail Kanaga Consulting est une solution intégrée et collaborative conçue pour piloter la performance individuelle et collective des équipes. Développé sur l'écosystème Google Apps Script avec stockage sur Google Sheets et intégration Google Drive/Docs, il garantit une fluidité totale entre l'auto-évaluation du collaborateur, la contribution d'évaluateurs pairs (secondaires) et l'arbitrage managérial final.")
    add_body("Ce dispositif assure la conformité déontologique, la traçabilité des compétences, l'alignement sur les standards de performance de l'entreprise et la génération instantanée de comptes-rendus contractuels au format PDF.")

    # 2. ACTEURS ET RÔLES
    add_h1("2. Acteurs & Rôles dans le Processus")
    add_body("Le système articule 4 rôles opérationnels distincts :")

    # Table for Roles
    t_roles = doc.add_table(rows=5, cols=2)
    t_roles.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers_roles = ["Rôle", "Responsabilités & Droits dans le Système"]
    col_widths_roles = [Inches(2.2), Inches(4.6)]

    for j, h in enumerate(headers_roles):
        cell = t_roles.cell(0, j)
        cell.width = col_widths_roles[j]
        set_cell_background(cell, "7C4C26")
        set_cell_margins(cell, 120, 120, 160, 160)
        p = cell.paragraphs[0]
        r = p.add_run(h)
        r.font.name = "Arial"
        r.font.bold = True
        r.font.size = Pt(10.5)
        r.font.color.rgb = RGBColor(255, 255, 255)

    roles_data = [
        ("Collaborateur (Évalué)", "Remplit son auto-évaluation (notation, commentaires, réalisations de la période, souhaits de formation et aspirations). Accède à l'historique de ses évaluations et télécharge ses comptes-rendus PDF."),
        ("Évaluateur(s) Secondaire(s)", "Pairs, chefs de missions ou collaborateurs mandatés. Fournissent une appréciation et une notation intermédiaire avant l'entretien final. Disposent de sauvegardes de brouillons indépendantes."),
        ("Évaluateur Principal (Manager)", "Initie l'évaluation, pilote l'entretien d'évaluation, consulte les avis du collaborateur et des secondaires en miroir, réalise la synthèse finale, fixe les objectifs SMART et clôture le dossier."),
        ("Administrateur RH / IT", "Administre le Form Builder dynamique (ajout de profils métiers, questions, pagination), personnalise les modèles d'e-mails automatiques, gère les habilitations et supervise les logs d'activité.")
    ]

    for i, (role, desc) in enumerate(roles_data):
        row = t_roles.rows[i+1]
        bg = "FBF9F6" if i % 2 == 0 else "FFFFFF"
        for j, val in enumerate([role, desc]):
            cell = row.cells[j]
            cell.width = col_widths_roles[j]
            set_cell_background(cell, bg)
            set_cell_margins(cell, 100, 100, 140, 140)
            p = cell.paragraphs[0]
            r = p.add_run(val)
            r.font.name = "Arial"
            r.font.size = Pt(10)
            if j == 0:
                r.font.bold = True
                r.font.color.rgb = primary_color
            else:
                r.font.color.rgb = dark_text

    # 3. WORKFLOW DE BOUT EN BOUT
    add_h1("3. Cycle de Vie & Workflow d'Évaluation de Bout en Bout")
    add_body("Le cycle d'évaluation suit un flux rigoureux composé de cinq phases séquentielles :")

    add_h2("Phase 1 : L'Initiation du Dossier")
    add_bullet("Déclencheur : ", "Le Manager ou l'Administrateur démarre le processus depuis le Portail Manager.")
    add_bullet("Paramètres saisis : ", "Période d'évaluation (ex: Année 2025), sélection des collaborateurs (support d'initiation en lot/batch), choix du profil de poste, sélection de l'évaluateur principal et des éventuels évaluateurs secondaires.")
    add_bullet("Statut du dossier : ", "Initiée.")
    add_bullet("Automatismes : ", "Notification par e-mail au collaborateur avec lien vers son auto-évaluation, alerte aux évaluateurs secondaires pour prise en compte, et e-mail de confirmation au manager.")

    add_h2("Phase 2 : L'Auto-évaluation du Collaborateur")
    add_bullet("Accès sécurisé : ", "Le collaborateur accède à son espace 'Mes Évaluations'. Les données générales de cadrage (profil, période, évaluateurs) sont verrouillées en lecture seule.")
    add_bullet("Saisie du collaborateur : ", "Évaluation des compétences fondamentales, des compétences spécifiques au profil métier, et renseignement de la conclusion (bilan des objectifs passés, points forts, axes d'amélioration, note globale, besoins de formation, aspirations professionnelles).")
    add_bullet("Sauvegarde intermédiaire : ", "Possibilité d'enregistrer des brouillons à tout moment sans notifier les autres acteurs.")
    add_bullet("Statut après soumission : ", "'Attente Évaluateur Secondaire' (si secondaires désignés) ou 'Attente Évaluateur Principal' (si aucun secondaire).")

    add_h2("Phase 3 : L'Évaluation Secondaire (Multi-évaluateurs)")
    add_bullet("Accès partagé et indépendant : ", "Chaque évaluateur secondaire accède au dossier et visualise les réponses et justifications de l'auto-évaluation du collaborateur.")
    add_bullet("Brouillon individuel : ", "Chaque évaluateur secondaire peut enregistrer son travail en brouillon sans impacter les autres évaluateurs secondaires.")
    add_bullet("Transition automatique ou manuelle : ", "Lorsque tous les secondaires ont soumis leur évaluation, le statut bascule automatiquement vers 'Attente Évaluateur Principal'. Le manager dispose en outre d'une option 'Passer au Principal →' pour forcer l'avancement si un avis tarde à être émis.")

    add_h2("Phase 4 : L'Évaluation Managériale & Entretien d'Évaluation")
    add_bullet("Vue comparative miroir 360° : ", "Dans son interface d'évaluation, le manager visualise côte à côte les notes et commentaires du collaborateur ainsi que l'ensemble des retours des évaluateurs secondaires.")
    add_bullet("Mode Pré-évaluation ('Pré-évaluer pour réunion') : ", "Le manager peut préparer en amont de l'entretien physique/visio ses notes et appréciations sans clôturer le dossier.")
    add_bullet("Option de retour aux secondaires : ", "Si des compléments sont requis, le manager peut renvoyer le dossier aux évaluateurs secondaires ('← Retourner aux Secondaires'). Ses propres réponses pré-saisies sont automatiquement conservées en brouillon.")
    add_bullet("Finalisation managériale : ", "Arbitrage sur les notes de compétences, validation des axes d'amélioration, définition conjointe des objectifs SMART futurs et attribution de la note globale définitive.")

    add_h2("Phase 5 : Clôture, Génération PDF & Archivage")
    add_bullet("Statut final : ", "Complétée.")
    add_bullet("Génération instantanée du PDF : ", "Déclenchement du script generateEvaluationPDF(). Un document Google Docs formaté est automatiquement créé, intégrant les tableaux récapitulatifs, la légende du barème et les colonnes comparatives (Collaborateur / Secondaires / Principal).")
    add_bullet("Archivage Google Drive : ", "Conversion automatique en PDF natif et stockage dans le dossier Drive dédié.")
    add_bullet("Notification finale : ", "Envoi automatique d'un e-mail à l'employé, au manager et aux secondaires avec bouton d'accès direct et de téléchargement du compte-rendu officiel.")

    # 4. ARCHITECTURE DU FORMULAIRE
    add_h1("4. Architecture du Formulaire d'Évaluation (L'Entonnoir en 4 Étapes)")
    add_body("Le formulaire d'évaluation a été conçu selon le principe d'un entonnoir dynamique structuré en 4 sections :")

    t_struct = doc.add_table(rows=5, cols=3)
    t_struct.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers_struct = ["Étape / Section", "Intitulé", "Contenu & Rôle Fonctionnel"]
    col_widths_struct = [Inches(1.2), Inches(2.2), Inches(3.4)]

    for j, h in enumerate(headers_struct):
        cell = t_struct.cell(0, j)
        cell.width = col_widths_struct[j]
        set_cell_background(cell, "7C4C26")
        set_cell_margins(cell, 120, 120, 160, 160)
        p = cell.paragraphs[0]
        r = p.add_run(h)
        r.font.name = "Arial"
        r.font.bold = True
        r.font.size = Pt(10)
        r.font.color.rgb = RGBColor(255, 255, 255)

    struct_data = [
        ("Étape 1", "Section 0 : Entonnoir de départ", "Informations générales : Période d'évaluation, employé(s) ciblé(s), profil de poste, désignation de l'évaluateur principal et des évaluateurs secondaires."),
        ("Étape 2", "Section 1 : Tronc Commun", "8 compétences fondamentales partagées par l'ensemble des collaborateurs de Kanaga Consulting :\n1. Professionnalisme et éthique\n2. Communication Orale et Écrite\n3. Travail d'équipe et Collaboration\n4. Organisation et Gestion du Temps\n5. Initiative et Proactivité\n6. Adaptabilité et Apprentissage Continu\n7. Compréhension du Contexte Local\n8. Capacité d'Innovation"),
        ("Étape 3", "Section 2 : Aiguillage Profil", "Compétences spécifiques au métier avec pagination dynamique (ex: Junior Secteur Privé, Assistant Secteur Public, Consultant, Comptable, Chef de Projet, Auditeur ou profil sur-mesure)."),
        ("Étape 4", "Section 3 : Conclusion Commune", "Synthèse globale et projection d'avenir :\n• Réalisation des objectifs de la période écoulée\n• Principaux points forts\n• Axes d'amélioration prioritaires\n• Appréciation de la performance globale\n• Besoins en formation\n• Objectifs SMART futurs\n• Aspirations professionnelles\n• Commentaires libres")
    ]

    for i, (st, tit, desc) in enumerate(struct_data):
        row = t_struct.rows[i+1]
        bg = "FBF9F6" if i % 2 == 0 else "FFFFFF"
        for j, val in enumerate([st, tit, desc]):
            cell = row.cells[j]
            cell.width = col_widths_struct[j]
            set_cell_background(cell, bg)
            set_cell_margins(cell, 100, 100, 140, 140)
            p = cell.paragraphs[0]
            r = p.add_run(val)
            r.font.name = "Arial"
            r.font.size = Pt(9.5)
            if j == 0:
                r.font.bold = True
                r.font.color.rgb = primary_color
            elif j == 1:
                r.font.bold = True
                r.font.color.rgb = dark_text
            else:
                r.font.color.rgb = dark_text

    # 5. BAREME D'EVALUATION
    add_h1("5. Le Barème Harmonisé d'Évaluation")
    add_body("Pour assurer l'équité, la rigueur et l'objectivité de la notation, un barème unique à 5 échelons (+ N/A) est appliqué sur l'ensemble de la plateforme :")

    t_scale = doc.add_table(rows=7, cols=3)
    t_scale.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers_scale = ["Niveau / Sigle", "Intitulé", "Définition & Critère d'Attribution"]
    col_widths_scale = [Inches(1.5), Inches(2.3), Inches(3.0)]

    for j, h in enumerate(headers_scale):
        cell = t_scale.cell(0, j)
        cell.width = col_widths_scale[j]
        set_cell_background(cell, "7C4C26")
        set_cell_margins(cell, 120, 120, 160, 160)
        p = cell.paragraphs[0]
        r = p.add_run(h)
        r.font.name = "Arial"
        r.font.bold = True
        r.font.size = Pt(10)
        r.font.color.rgb = RGBColor(255, 255, 255)

    scale_data = [
        ("5 - EXCEPTI", "Performance Exceptionnelle", "Dépasse constamment les attentes et constitue une référence pour l'organisation.", "1E7E34"),
        ("4 - SUPERIE", "Performance Supérieure", "Dépasse régulièrement les attentes dans les domaines clés de la mission.", "0C63E4"),
        ("3 - CONFORM", "Performance Conforme aux Attentes", "Atteint les objectifs et les standards de performance attendus pour le poste.", "0F5132"),
        ("2 - AMELIOR", "Performance à Améliorer", "N'atteint pas toujours les objectifs et nécessite une amélioration dans certains domaines.", "856404"),
        ("1 - INSATIS", "Performance Insatisfaisante", "N'atteint pas les objectifs de manière significative ; nécessite une amélioration immédiate.", "842029"),
        ("NONAPPL", "N/A Non Applicable", "Non applicable au poste ou non évaluable sur la période de référence.", "555555")
    ]

    for i, (lvl, tit, desc, hex_color) in enumerate(scale_data):
        row = t_scale.rows[i+1]
        bg = "FBF9F6" if i % 2 == 0 else "FFFFFF"
        for j, val in enumerate([lvl, tit, desc]):
            cell = row.cells[j]
            cell.width = col_widths_scale[j]
            set_cell_background(cell, bg)
            set_cell_margins(cell, 90, 90, 130, 130)
            p = cell.paragraphs[0]
            r = p.add_run(val)
            r.font.name = "Arial"
            r.font.size = Pt(9.5)
            if j == 0:
                r.font.bold = True
                r.font.color.rgb = RGBColor.from_string(hex_color)
            elif j == 1:
                r.font.bold = True
                r.font.color.rgb = dark_text
            else:
                r.font.color.rgb = dark_text

    # 6. POINTS FONDAMENTAUX DES FONCTIONNALITES
    add_h1("6. Points Fondamentaux & Fonctionnalités Techniques Clés")

    add_h2("1. Form Builder Administrateur & Modularité")
    add_body("Les formulaires d'évaluation sont pilotés dynamiquement via la feuille Google Sheets 'EvaluationQuestions'. L'administrateur peut à tout moment créer de nouveaux profils métiers, ajuster les questions, modifier les descriptions, ordonner les sections et réutiliser des pages existantes grâce à la modale de duplication de pages.")

    add_h2("2. Moteur de Notifications & Templates Paramétrables")
    add_body("La plateforme intègre un centre de gestion des e-mails automatiques configuré sur 5 événements majeurs (initiation collaborateur, invitation secondaires, alerte auto-évaluation soumise, alerte retours secondaires complétés, finalisation avec lien PDF). L'administrateur peut personnaliser l'objet et le corps de chaque message avec des balises dynamiques ({nom_employe}, {periode}, {profil}, {evaluateur_principal}, {evaluateurs_secondaires}, {lien_portail}, {bloc_pdf}) et envoyer des e-mails de test.")

    add_h2("3. Gestion Avancée des Brouillons & Sécurité des Données")
    add_body("Afin d'éviter tout conflit de saisie, les brouillons de chaque acteur (collaborateur, évaluateurs secondaires, manager) sont isolés dans des colonnes JSON dédiées. Si un manager réassigne un dossier aux secondaires, ses réponses préalables restent intactes en brouillon.")

    add_h2("4. Générateur Automatique de Compte-Rendu PDF")
    add_body("À la clôture de l'évaluation, le système génère un document officiel au format PDF sous Google Drive via DocumentApp. Ce document récapitule les métadonnées de l'entretien, le barème d'évaluation, le tableau comparatif complet (Employé / Secondaires / Principal) et la synthèse globale. Le lien Drive est archivé dans le portail et inséré dans les e-mails de clôture.")

    add_h2("5. Traçabilité, Sécurité & Logs")
    add_body("Toutes les actions stratégiques (initiation, soumission, brouillons, clôture) sont tracées en temps réel par la fonction logEvent(). Chaque action est horodatée avec l'identifiant de l'utilisateur, l'action réalisée et son statut, puis enregistrée dans Google Cloud Logging et dans l'onglet 'Logs' du tableur.")

    # 7. FEUILLE DE ROUTE & AXES D'AMÉLIORATION
    add_h1("7. Feuille de Route & Axes d'Amélioration Technique")
    add_body("Pour garantir la pérennité, la rapidité et la robustesse de l'outil lors des futures montées en charge, six chantiers majeurs sont identifiés :")

    t_roadmap = doc.add_table(rows=7, cols=3)
    t_roadmap.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers_roadmap = ["Chantier", "Objectif & Action Recommandée", "Priorité"]
    col_widths_roadmap = [Inches(1.8), Inches(3.9), Inches(1.3)]

    for j, h in enumerate(headers_roadmap):
        cell = t_roadmap.cell(0, j)
        cell.width = col_widths_roadmap[j]
        set_cell_background(cell, "7C4C26")
        set_cell_margins(cell, 120, 120, 160, 160)
        p = cell.paragraphs[0]
        r = p.add_run(h)
        r.font.name = "Arial"
        r.font.bold = True
        r.font.size = Pt(10)
        r.font.color.rgb = RGBColor(255, 255, 255)

    roadmap_data = [
        ("1. Modularisation du Code", "Scinder Code.js (3 900+ lignes) et javascript.html (7 000+ lignes) en services spécialisés (AuthService, EvaluationService, ReminderService, QuestionsService).", "Haute", "B71C1C"),
        ("2. Cache & Performance", "Implémenter CacheService pour stocker l'arborescence des questions et profils métiers pendant 30 min. Réduire de 80% les accès répétitifs à Google Sheets.", "Haute", "B71C1C"),
        ("3. Concurrence & Verrous", "Encapsuler les sauvegardes et clôtures critiques dans LockService.getScriptLock() afin d'éliminer les risques de collision d'écriture simultanée.", "Haute", "B71C1C"),
        ("4. Sécurité & Contrôle d'Accès", "Systématiser la vérification des permissions par évaluation (prévention IDOR) et basculer les identifiants sensibles vers PropertiesService.", "Moyenne", "E65100"),
        ("5. Résilience Offline (UX)", "Intégrer une double persistance des réponses sur localStorage / IndexedDB pour éviter toute perte de saisie en cas de micro-coupure réseau.", "Moyenne", "E65100"),
        ("6. Automatisation CI/CD", "Mettre en place un workflow GitHub Actions pour exécuter les tests et déployer automatiquement sur la branche main avec clasp.", "Évolution", "2E7D32")
    ]

    for i, (ch, desc, prio, hex_col) in enumerate(roadmap_data):
        row = t_roadmap.rows[i+1]
        bg = "FBF9F6" if i % 2 == 0 else "FFFFFF"
        for j, val in enumerate([ch, desc, prio]):
            cell = row.cells[j]
            cell.width = col_widths_roadmap[j]
            set_cell_background(cell, bg)
            set_cell_margins(cell, 90, 90, 130, 130)
            p = cell.paragraphs[0]
            r = p.add_run(val)
            r.font.name = "Arial"
            r.font.size = Pt(9.5)
            if j == 0:
                r.font.bold = True
                r.font.color.rgb = primary_color
            elif j == 2:
                r.font.bold = True
                r.font.color.rgb = RGBColor.from_string(hex_col)
            else:
                r.font.color.rgb = dark_text

    # 8. GUIDE DES FUTURES MISES À JOUR & DÉPLOIEMENT
    add_h1("8. Protocole Opérationnel des Futures Mises à Jour")
    add_body("Pour toute mise à jour ultérieure, l'équipe technique doit suivre la procédure standard suivante :")
    add_bullet("1. Récupération : ", "git pull origin main et npx clasp pull pour synchroniser le code.")
    add_bullet("2. Développement : ", "Édition du code dans le respect de la charte visuelle Kanaga et des conventions de nommage.")
    add_bullet("3. Sauvegarde Git : ", "git add -A puis git commit -m 'feat: ...' et git push origin main.")
    add_bullet("4. Déploiement Permanent : ", "Exécuter 'node deploy.js' pour pousser le code et mettre à jour le déploiement fixe sans modifier l'URL publique des utilisateurs.")
    add_bullet("5. URL Permanente Invariable : ", "https://script.google.com/macros/s/AKfycbxeoDfu8Uh9plmQXjud3N1cXUBmSAaIpgdrXQxgWEnp1jst3K3N2puD1pF3zl1XYsRntA/exec")

    output_path = r"c:\Users\Mahamane\Documents\Project\New_Kanaga-Portail-evaluation\Rapport_Processus_Evaluation_Kanaga.docx"
    doc.save(output_path)
    print("Document successfully created at:", output_path)

if __name__ == '__main__':
    create_report()

